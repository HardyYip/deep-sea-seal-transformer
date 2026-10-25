"""Build the Chinese completion report from verified, frozen experiment records.

Run with the Codex bundled Python. No training, model inference or source edits.
"""
import json
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/final_report'
A = json.loads((OUT/'analysis.json').read_text(encoding='utf-8'))
DOCX = OUT/'深海柔性密封性能预测项目结题报告.docx'
MD = OUT/'深海柔性密封性能预测项目结题报告.md'
D = Document()
section = D.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2.0), Cm(1.9)
section.left_margin, section.right_margin = Cm(2.15), Cm(2.15)
section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
CONTENT_W = 16.7
markdown = []


def set_font(style, latin='Times New Roman', east='宋体', size=10.5, bold=False):
    style.font.name=latin; style.font.size=Pt(size); style.font.bold=bold
    style.font.italic=False
    style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for k,v in [('ascii',latin),('hAnsi',latin),('eastAsia',east)]:rf.set(qn('w:'+k),v)


for name in ['Normal','Body Text']:
    set_font(D.styles[name]); pf=D.styles[name].paragraph_format
    pf.line_spacing=1.35; pf.space_after=Pt(6); pf.widow_control=True
for name,size in [('Title',23),('Subtitle',12),('Heading 1',16),('Heading 2',12.5),('Heading 3',11)]:
    set_font(D.styles[name],east='黑体',size=size,bold=name!='Subtitle')
    pf=D.styles[name].paragraph_format; pf.space_before=Pt(12 if name.startswith('Heading') else 0)
    pf.space_after=Pt(8); pf.keep_with_next=True
    # Remove themed title borders if a renderer supplies them.
    pp=D.styles[name].element.find(qn('w:pPr'))
    if pp is not None:
        b=pp.find(qn('w:pBdr'))
        if b is not None:pp.remove(b)
set_font(D.styles['Caption'],east='宋体',size=9)
D.styles['Caption'].paragraph_format.space_after=Pt(7)
D.core_properties.title='深海柔性密封性能预测项目结题报告'
D.core_properties.subject='最终 Transformer 方案的代码分析 数据处理 建模过程 冻结测试结果与方案比较'
D.core_properties.author=''
D.core_properties.keywords='柔性密封 Transformer 多输出回归 数据审计 冻结测试'
footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fr=footer.add_run(); fr.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); fr._r.addnext(field)


def p(text, bold_start=None):
    para=D.add_paragraph()
    if bold_start and text.startswith(bold_start):
        para.add_run(bold_start).bold=True; para.add_run(text[len(bold_start):])
    else:para.add_run(text)
    markdown.append(text+'\n')
    return para


def h(text, level=1):
    if level>1:
        text=re.sub(r'^\d+\.\d+\s+', '', text)
    text=text.replace('p_ocean','海洋压力')
    D.add_heading(text,level=level); markdown.append('#'*(level+1)+' '+text+'\n')


def page(title):
    D.add_page_break(); markdown.append('\n<!-- page break -->\n'); h(title)


def table(headers, rows, widths=None, caption=None, left_cols=(0,)):
    if caption:
        cap=D.add_paragraph(caption,style='Caption');cap.paragraph_format.keep_with_next=True
        markdown.append(caption+'\n')
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [CONTENT_W/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for c,w in zip(t.rows[0].cells,widths):c.width=Cm(w)
    trpr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader');trpr.append(rep)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for vals in rows:
        cells=t.add_row().cells
        for c,w in zip(cells,widths):c.width=Cm(w)
    for i,vals in enumerate([headers]+list(rows)):
        row=t.rows[i]; rr=row._tr.get_or_add_trPr(); avoid=OxmlElement('w:cantSplit');rr.append(avoid)
        for j,(cell,text) in enumerate(zip(row.cells,vals)):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for side,n in [('top',72),('bottom',72),('left',90),('right',90)]:
                v=OxmlElement('w:'+side);v.set(qn('w:w'),str(n));v.set(qn('w:type'),'dxa');margins.append(v)
            tcpr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE7EE' if i==0 else ('F5F7F8' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            para=cell.paragraphs[0];para.paragraph_format.line_spacing=1.15
            para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0)
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            r=para.add_run(str(text));r.font.size=Pt(9);r.bold=i==0
    markdown.append('| '+' | '.join(map(str,headers))+' |\n| '+' | '.join(['---']*len(headers))+' |')
    markdown.extend('| '+' | '.join(str(c).replace('\n','<br>') for c in row)+' |' for row in rows)
    markdown.append('')
    D.add_paragraph().paragraph_format.space_after=Pt(0)


def fig(name,caption,width=16.3):
    para=D.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True;para.paragraph_format.space_after=Pt(3)
    r=para.add_run();r.add_picture(str(OUT/'figures'/f'{name}.png'),width=Cm(width))
    inline=r._r.find('.//'+qn('wp:docPr'))
    if inline is not None:inline.set('descr',caption)
    D.add_paragraph(caption,style='Caption')
    markdown.append(f'![{caption}](figures/{name}.png)\n')


def code(text):
    para=D.add_paragraph();para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(8)
    para.paragraph_format.line_spacing=1.05
    r=para.add_run(text);r.font.name='Consolas';r.font.size=Pt(8.2)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    markdown.append('```'+('powershell' if text.startswith('& ') else 'python')+'\n'+text+'\n```\n')


def eq(text):
    para=D.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMathPara');om=OxmlElement('m:oMath')
    def run(parent,value):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value
        t.set(qn('xml:space'),'preserve');r.append(t);parent.append(r)
    # Convert deliberately limited subscript/superscript notation to native OMML.
    pattern=re.compile(r'([A-Za-zℝμσθŷ]+)(?:_\{([^}]+)\})?(?:\^\{([^}]+)\})?')
    cursor=0
    for match in pattern.finditer(text):
        run(om,text[cursor:match.start()]); base,sub,sup=match.groups()
        if sub is not None or sup is not None:
            item=OxmlElement('m:sSubSup' if sub is not None and sup is not None else ('m:sSub' if sub is not None else 'm:sSup'))
            e=OxmlElement('m:e');run(e,base);item.append(e)
            if sub is not None:
                e=OxmlElement('m:sub');run(e,sub);item.append(e)
            if sup is not None:
                e=OxmlElement('m:sup');run(e,sup);item.append(e)
            om.append(item)
        else:run(om,base)
        cursor=match.end()
    run(om,text[cursor:]);math.append(om);para._p.append(math)
    markdown.append('$$\n'+text+'\n$$\n')


def num(v, digits=5):
    return f'{v:.{digits}f}' if abs(v)<.1 else f'{v:.4f}'


# 1
D.add_paragraph('深海柔性密封性能预测\n项目结题报告',style='Title')
D.add_paragraph('基于节点特征与全局压力融合的双分组 Transformer 回归',style='Subtitle')
p('报告日期  2026 年 10 月 13 日\n最终数据版本  v2    最终实验  run_004')
h('摘要')
p('本项目建立了面向深海柔性密封工况的监督回归流程，分别为 positive 和 negative 两个数据分组训练独立 Transformer，同时预测 ff、f、q，保留原始正负号。输入由 221 个节点的三通道特征及两个全局压力组成；用户确认的摩擦力由原尺度 ff 与 f 相加得到。')
p('工作覆盖原始数据审计、来源核对、异常与关联处理、固定清单划分、训练集标准化、GPU 训练、验证集早停、独立测试及无目标列推理。原始数据共 2093 条，最终实验暂时排除 negative 的 3 条指定记录，使用 2090 条。positive 与 negative 测试集分别为 153 条和 149 条。')
table(['分组','ff 测试 R²','f 测试 R²','q 测试 R²'],[
    [g,*[f"{A['groups'][g]['test']['transformer'][t]['R2']:.4f}" for t in ('ff','f','q')]] for g in ('positive','negative')],
    [4.1,4.2,4.2,4.2],caption='表 1  最终冻结测试结果概览')
p('与训练集目标均值基线相比，六组目标的测试 MAE 降低约 91.09% 至 97.30%。结果支持模型在现有压力范围及现有输入条件下的工况回归能力；p_ocean=0 的边界工况仍存在显著局部误差。当前实验尚不足以证明未见压力等级的泛化、范围外外推或实际工程安全裕度。')
p('关键词  柔性密封  多输出回归  Transformer  节点特征  数据溯源  独立测试')
h('报告内容')
p('项目目标与数据定义；数据审计和清洗；划分与标准化；模型结构与数学表达；关键代码分析；训练与开发实验；最终结果与边界误差；方案优劣比较；可复现性、局限与结题成果。')

# 2
page('1 项目目标与问题定义')
h('1.1 预测任务',2)
p('一条样本表示一个数据分组下的一组压力工况，221 个节点共同描述该工况的节点状态。研究目标是学习从节点状态及压力到三个连续输出的映射，而非把节点拆分为独立训练样本。原表没有明确时间列，本项目属于工况回归。')
eq('g_{θ} : (X ∈ ℝ^{221×3}, P ∈ ℝ^{2}) → ŷ ∈ ℝ^{3}')
table(['变量或标识','模型角色','维度或单位','说明'],[
    ['p_system  p_ocean','全局输入','2 维  MPa','保留原始数值'],
    ['x1…x221','节点输入','221 维','确切定义与单位待确认'],
    ['y1…y221','节点输入','221 维','不是监督目标 Y'],
    ['cpress1…cpress221','节点输入','221 维','确切定义与单位待确认'],
    ['ff  f  q','监督输出','N  N  mm³/s','顺序固定为 ff f q'],
    ['sample_id  来源  split','溯源或分区','不进入特征','用于对齐和复核']
],[4.3,2.6,3.4,6.4],caption='表 2  输入输出和元信息定义',left_cols=(0,3))
h('1.2 双分组建模与符号约定',2)
p('positive、negative 沿用原始分组名称。用户转述其与压力正负有关，但压力列本身均为非负数，具体约定仍未完全明确，因此不将 negative 压力取负，也不将两个分组解释为分类标签或直接认定为运动方向。')
p('两个模型结构相同、权重独立，每个模型都输出三个目标。ff 与 f 的更细力学定义尚需补充；已确认的关系是摩擦力等于 ff+f。两个目标均不作为另一个目标的输入。线性输出层允许正、负和零预测，训练及评价均保留原符号。')
h('1.3 项目成果的适用条件',2)
p('预测时必须提供与训练一致的全部节点特征及两个压力，仅给出压力无法运行当前模型。若节点状态依赖先完成昂贵仿真，则本模型的定位是具备这些输入条件下的代理回归，现阶段没有证据表明可以独立替代完整仿真流程。')

# 3
page('2 数据来源 审计与清洗')
h('2.1 数据池及结构核验',2)
p('实际只读来源为项目 src/positive 和 src/negative。每组包含 5 至 70 MPa、间隔 5 MPa 的 14 个压力目录；只纳入文件名匹配 *_training.csv 的文件。重新统计 positive 为 1049 条、negative 为 1044 条，共 2093 条，均为 668 列，其中输入 665 列、目标 3 列。')
p('整理脚本校验列名、重复表头、x/y/cpress 的 1 至 221 连续编号、数值解析、NaN/Inf、目录压力一致性、重复工况及相同完整输入。记录 sample_id、来源相对路径和原 CSV 行号，并保存哈希。数值审计未发现缺失、非数值或非有限数值；并不以此推断物理有效性。')
h('2.2 来源对照与异常处置',2)
p('已有证据复核读取 2095 个 XLSX，其中 2094 个具备 Result 及四张分布表，另 1 个为空白；配套普通 CSV 和 XLSX 只用于核对，不重复加入训练池。源表一致性支持合并过程正确，但导出文件缺少可核实的残差、容差和成功终止标记。')
table(['工况  MPa','证据或处置','最终状态'],[
    ['negative  15 0','ff = −1854.975068 N  源表一致','用户指定暂时排除'],
    ['negative  20 0','ff = −405.437850 N  源表一致','用户指定暂时排除'],
    ['negative  25 0','按用户指定追加排除','用户指定暂时排除'],
    ['negative  65 60…65','11 条共享全部 663 个节点值','保守关联成组 全部 train']
],[4.1,7.5,5.1],caption='表 3  最终实验中的重点数据处理',left_cols=(0,1,2))
p('三条排除记录在原始文件和完整合并表中继续保留，排除不代表永久无效。清洗没有取绝对值、截断、平滑、插值或自动按 3σ/IQR 删除；未补齐池外异常或空白工况。11 条共享节点记录的 p_ocean 及目标不同，因此不按重复整行删除，也不将目标平均。')
h('2.3 数据资格与证据边界',2)
p('实验包将 2090 条选用记录标为 provisional，usage 为 user_authorized_experiment，physical_review_status 为 not_fully_verified。用户补充“每个 node 之间是收敛的”，该信息属于用户说明，未定义完整工况的求解收敛判据。结题成果因此是已完成的数值建模实验，物理审核状态另行保留。')

# 4
page('3 数据划分与标准化')
h('3.1 固定工况清单与关联隔离',2)
p('初始方案使用 seed=42，分别在两个分组内按 p_system 分层，以完整工况及关联组为单位分配约 70% train、15% validation 和 15% test。节点不拆分。整数分配、共享节点归组、暂时排除及边界定向重分层使最终比例略有变化。')
fig('06_split_counts','图 1  v2 数据包实际分区数量与比例',width=15.8)
p('positive 为 738/158/153，共 1049 条；negative 为 740/152/149，共 1041 条。negative 的 11 条共享节点全部归 train；p_ocean=0 的 9 条非排除边界记录固定为 train 5、validation 2、test 2，其中 50、55、70 MPa 从原 validation 调至 train。调整发生于最终冻结评价之前，不在测试后重新抽样或补足比例。')
h('3.2 训练集拟合的三类缩放',2)
p('标准化只使用 train。节点按 x、y、cpress 三种通道，跨训练样本与全部节点统计均值和标准差；两个压力分别统计；三个目标分别统计。常量标准差替换为 1。validation 与 test 只使用保存参数做变换。')
eq('z = (v − μ_{tr}) / σ_{tr}     ŷ_{orig} = ŷ_{std} ⊙ σ_{y} + μ_{y}')
p('标准化用于数值优化，不改变结果报告的原始单位。元信息保存拟合 sample_id、通道与目标顺序。最终 positive 和 negative 缩放拟合 ID 分别为 738 个和 740 个，核对后与各自 train 完全一致。')
h('3.3 评价分工',2)
p('train 用于优化权重；validation 用于选择超参数与最佳 checkpoint；test 在方案冻结后单独评价。按压力层内留出工况只能检验现有压力范围内的条件预测，不能证明未见 p_system 压力等级或范围外工况的外推能力。')

# 5
page('4 最终模型结构')
fig('01_architecture','图 2  节点注意力与全局压力融合的模型结构')
table(['模块','输入与输出','参数量'],[
    ['节点线性映射','[B 221 3] → [B 221 64]','256'],
    ['2 层 Transformer Encoder','[B 221 64] → [B 221 64]','66 944'],
    ['节点均值池化','[B 221 64] → [B 64]','0'],
    ['压力 Linear 与 GELU','[B 2] → [B 64]','192'],
    ['融合回归头','[B 128] → [B 64] → [B 3]','8 451'],
    ['单个分组模型合计','双模型参数互不共享','75 843']
],[5.8,7.2,3.7],caption='表 4  各模块维度和可训练参数',left_cols=(0,1))
p('每个节点的 [xi, yi, cpressi] 经共享线性层投影为 64 维，再进入两层编码器。编码器使用 4 个注意力头，每头维度为 16，前馈隐藏维度为 128，激活为 GELU，dropout 为 0.1。各层包含注意力、前馈网络、残差连接与层归一化；代码采用 PyTorch 编码器层的默认后归一化行为。')
p('节点表示沿节点维求均值得到全局节点摘要。两个压力经独立 Linear 与 GELU 映射为 64 维，随后与节点摘要拼接，再由 128→64→3 回归头输出标准化目标。压力在节点编码完成后融合，当前注意力权重不直接由压力分支条件化。')

# 6
page('5 建模表达与结构选择依据')
h('5.1 节点间信息交互',2)
eq('H^{0} = XW_{n} + b_{n}')
eq('Attention(Q,K,V) = softmax(QK^{T} / √d_{k})V')
eq('h_{n} = meanᵢ(H_{i}^{2})     h_{p} = GELU(PW_{p} + b_{p})')
eq('ŷ = Linear_{3}(Dropout(GELU(Linear_{64}([h_{n} ; h_{p}]))))')
p('注意力让节点表示结合其他节点的信息，从而学习节点状态与整体输出之间的统计关系。多头机制为不同特征子空间提供并行交互；两层编码器控制了模型规模。这里的注意力是可学习的统计运算，不能直接解释为真实接触力传递或已识别的力学因果机制。[R1]')
h('5.2 无位置编码与节点重排不变性',2)
p('节点编号的物理对应尚未完全确认，代码没有引入时间或编号位置编码。共享节点映射与无位置编码自注意力对完整节点三元组的排列具有等变性，配合均值池化后，推理输出对节点三元组整体重排近似不变，差异仅可能来自浮点运算。训练时 dropout 会引入随机扰动。[R2]')
p('该设计保留每个节点三通道的配对关系，同时避免把未知编号误当作时间顺序。代价是模型不能仅凭节点编号区分物理位置或恢复明确拓扑。如果后续确认 x、y 为坐标，它们可承载部分几何信息，但仍不能替代已经核实的邻接关系。当前实现还固定要求 221 个节点，不支持任意节点数或缺失节点掩码。')
h('5.3 双模型与联合三目标',2)
p('分组独立建模允许两组分别学习不同映射并采用不同学习率，避免在分组语义尚不充分明确时强行共享规律。代价是每组仅约千条样本，不能利用跨组共享数据提高统计效率。')
p('三个目标共享编码器和回归头的中间表示，减少三套独立特征提取器的重复参数。标准化后等权优化，使大数值的 f 不直接支配损失，但并不等于三个目标具有相同工程重要性，也没有显式强制摩擦力、泄漏或其他物理守恒约束。')

# 7
page('6 关键代码分析')
h('6.1 数据加载与正确节点组织',2)
p('models/data.py 的 load_group 以 sample_id 对齐清单和原值合并表，检查 ID 唯一性、集合一致性、工况及来源匹配，并重新验证工况、关联组、完整输入和节点块不跨分区。实验数据还须满足用途标识及文件指纹检查。')
code('nodes = np.stack([\n    np.array([[float(r[f"{c}{i}"]) for i in range(1, 222)]\n              for r in records])\n    for c in ("x", "y", "cpress")\n], axis=-1)')
p('原 CSV 按 x 块、y 块、cpress 块存储，逐块提取后沿最后一维堆叠，才得到 [N,221,3]。直接将原始 663 列 reshape 会把不同节点与通道混配。目标始终按 ff、f、q 读取，来源字段、split 和目标均不混入输入。')
h('6.2 缩放和前向计算',2)
code('scaler = fit_scaler(partitions["train"])\nencoded = self.encoder(self.node_projection(nodes)).mean(dim=1)\npressure = self.pressure_projection(global_features)\nprediction = self.head(torch.cat((encoded, pressure), dim=-1))')
p('fit_scaler 明确拒绝非 train 分区，transform 输出 float32，inverse_targets 恢复原尺度。模型前向检查节点及压力张量形状。Transformer 克隆层后，代码显式对注意力与线性权重做独立 Xavier 初始化，避免各编码层保持相同初值。')
table(['模块或入口','职责'],[
    ['prepare_data.py  review_data_evidence.py','合并审计及来源证据复核'],
    ['finalize_cleaning_review.py  prepare_training_bundle.py','固定清洗决定 关联划分 版本化实验包'],
    ['models/data.py  seal_transformer.py','清单对齐 缩放 指标与模型前向'],
    ['train_models.py  models/runtime.py','训练与独立测试入口 GPU 实际运算预检'],
    ['predict.py','无目标列输入 原尺度预测 摩擦力相加']
],[8.6,8.1],caption='表 5  核心代码职责',left_cols=(0,1))

# 8
page('7 训练目标与执行流程')
h('7.1 损失函数和优化策略',2)
eq('L = (1 / 3B) Σᵢ Σⱼ (ŷ_{ij} − t_{ij})²')
p('式中 t 与 ŷ 为标准化后的目标和预测，i 为 batch 内样本，j 依次对应 ff、f、q。')
p('训练使用标准化三个目标的平均 MSE，优化器为 AdamW，weight_decay=10⁻⁴；每个 batch 反向传播后按范数 1.0 裁剪梯度。非有限训练损失、验证损失或梯度触发错误。AdamW 将权重衰减与自适应梯度更新解耦，实际泛化效果仍由本项目验证结果判断。[R3]')
table(['训练设置','positive','negative'],[
    ['学习率','0.0003','0.001'],['编码器层数  注意力头','2 层  4 头','2 层  4 头'],
    ['d_model  前馈维度','64  128','64  128'],['dropout','0.1','0.1'],
    ['batch_size  seed','32  42','32  42'],['最大轮次  patience','100  15','100  15'],
    ['优化器  weight_decay','AdamW  0.0001','AdamW  0.0001'],['梯度裁剪范数','1.0','1.0']
],[7.1,4.8,4.8],caption='表 6  最终独立配置',left_cols=(0,))
h('7.2 验证集选择最佳权重',2)
p('每轮训练对 train 进行打乱抽样；validation 在 eval 与 no_grad 模式下完整预测。验证标准化平均 MSE 严格下降时保存 best.pt，连续 15 轮未改善则早停。训练结束加载最佳权重，输出原尺度 validation 指标及逐样本预测。没有使用 test 选择轮次或模型。')
h('7.3 GPU 环境与运行保护',2)
p('最终记录的运行环境为 Python 3.11、PyTorch 2.13.0+cu130、CUDA、RTX 5070 Laptop GPU。运行预检执行实际 Transformer 前向和反向，而不只判断 CUDA 可见。旧 pytorch 环境的内核错误不代表现有 pytorch2 不可用。')
p('训练输出目录必须是新目录，数据包不同内容拒绝覆盖已有版本。checkpoint 与数据指纹随元信息保存；独立 test 入口核对训练完成状态和指纹，并拒绝覆盖已有 test_metrics.json。这种机制有助于保持评价纪律，但无法从技术上阻止人工复制测试集或修改程序。')

# 9
page('8 开发实验与最终方案确定')
h('8.1 同版本开发实验比较',2)
p('run_001、run_002、run_003 保存的各组数据及清单指纹一致，分别比较两层 lr=0.001、三层 lr=0.001、两层 lr=0.0003。三次实验均使用 seed=42，历史记录的 test_evaluated 为 false，下表仅比较 validation。')
rows=[]
for group in ('positive','negative'):
    for d in A['development_runs']:
        if d['group']!=group:continue
        rows.append([group,d['run'].replace('run_',''),str(d['config']['model']['num_layers']),
                     str(d['config']['learning_rate']),f"{d['best_validation_mse']:.6f}",
                     *[f"{d['validation'][t]['R2']:.4f}" for t in ('ff','f','q')]])
table(['分组','运行','层数','lr','最佳验证 MSE','ff R²','f R²','q R²'],rows,
      [2.6,1.1,1.1,1.9,3.1,2.3,2.3,2.3],caption='表 7  旧数据版本内的开发实验',left_cols=(0,))
p('positive 的两层 lr=0.0003 得到最低标准化验证 MSE=0.006530，优于两层 lr=0.001 的 0.007226 和三层的 0.009389，支持最终采用较小学习率与两层结构。不同目标可能存在权衡，单个目标 MAE 最低不必对应联合损失最低。')
p('negative 的三层方案联合验证 MSE 略低于两层，但差值很小，同时 f 与 q 的单目标结果退化。两层 lr=0.001 在较低结构复杂度下保留了更平衡的结果，最终方案选择它；现有单次实验不能证明这一选择在所有随机种子下最优。')
h('8.2 数据修订与模型改进的区分',2)
p('run_004 的 negative 使用 v2：追加暂时排除 (25,0)，并将 3 条边界 validation 记录移入 train。validation 样本集合随之改变，不能将 run_004 的 ff 验证改善全部归因于学习率或模型结构，也不能把它与旧数据结果作为严格消融。positive 实际选用样本及划分未变，其最终验证结果与 run_003 在浮点精度范围内一致。')
p('最终方案因此是依据既有验证记录和数据处置形成的工程选择，而非已完成多随机种子统计检验或全面模型搜索的结论。最终 test 只在 run_004 冻结后评价。')

# 10
page('9 训练曲线与停止过程')
fig('02_learning_curves','图 3  最终训练与验证标准化 MSE 曲线  纵轴采用对数刻度')
table(['分组','最佳 epoch','停止 epoch','最佳验证 MSE','记录耗时  s'],[
    [g,A['groups'][g]['best_epoch'],A['groups'][g]['epochs'],f"{A['groups'][g]['best_validation_mse']:.6f}",f"{A['groups'][g]['elapsed_seconds']:.2f}"]
    for g in ('positive','negative')],[3.3,3.2,3.2,4.0,3.0],caption='表 8  训练终止与最佳 checkpoint',left_cols=(0,))
p('positive 最佳轮次为 79，连续 15 轮未改善后在 94 轮结束；negative 最佳轮次为 18，在 33 轮结束。最终使用最佳轮次权重，不能把停止轮次等同于实际交付 checkpoint。')
p('positive 首轮验证 MSE 约为 0.319072，最佳值为 0.006530；negative 由 0.037656 降至 0.011112。曲线展示了联合标准化损失下降，但并不证明每个局部工况均拟合良好。编码器优化的收敛也不能替代源仿真求解器的收敛证据。')
p('训练曲线的 train 值是在启用 dropout 且权重逐 batch 更新时累计得到，validation 值在 eval 模式下、该轮结束后计算。两条曲线的统计条件不同，因此 validation MSE 低于 train MSE 并不自动意味着泄漏。反过来，仅凭曲线形状也不能证明完全没有过拟合。')
p('元信息记录耗时约为 positive 17.10 s、negative 6.14 s，计时从训练主体开始，未覆盖完整数据审核、预检、测试和部署流程，也没有形成可与仿真时长公平比较的推理测速。本报告不据此宣称工程加速倍数。')
h('9.1 稳定性证据的限度',2)
p('当前每种配置只保存单个随机种子实验。固定 seed 有助于复现划分和随机初始化，但代码未启用完整确定性算法控制，GPU 运算仍可能存在微小数值差异。跨种子波动、置信区间及最坏工况鲁棒性尚未形成系统评价。')

# 11
page('10 冻结验证与测试成绩')
h('10.1 指标定义',2)
eq('MAE = mean(|ŷ − y|)     RMSE = √mean((ŷ − y)²)')
eq('R² = 1 − Σ(ŷ − y)² / Σ(y − mean(y))²')
p('所有指标都在目标逆标准化后计算，ff、f 的误差单位为 N，q 为 mm³/s。R² 衡量相对样本均值的解释程度，不是预测准确率。不同分组目标方差不同，不能只按 R² 横向判断绝对误差；不同单位也不直接相加为物理指标。')
rows=[]
for g in ('positive','negative'):
    for t in ('ff','f','q'):
        d=A['groups'][g];s=d['test']['transformer'][t]
        rows.append([g,t,num(d['validation'][t]['MAE']),num(s['MAE']),num(s['RMSE']),f"{s['R2']:.4f}"])
table(['分组','目标','验证 MAE','测试 MAE','测试 RMSE','测试 R²'],rows,[2.9,1.3,3.1,3.1,3.1,3.2],
      caption='表 9  最终原尺度验证与独立测试指标',left_cols=(0,))
p('positive 的 f、q 测试 R² 分别为 0.9960、0.9910；ff 为 0.9754，相比验证 ff R²=0.9881 有一定下降。negative 测试三目标 R² 均超过 0.9930，但 ff 测试 RMSE=0.03321 N 低于验证的 0.06439 N，说明少量边界样本对不同分区的影响明显。')
h('10.2 与训练均值基线的实测比较',2)
rows=[]
for g in ('positive','negative'):
    for t in ('ff','f','q'):
        d=A['groups'][g];b=d['test']['training_mean_baseline'][t]
        rows.append([g,t,num(b['MAE']),num(d['test']['transformer'][t]['MAE']),f"{d['test_mae_reduction_pct'][t]:.2f}%"])
table(['分组','目标','基线测试 MAE','模型测试 MAE','MAE 降低'],rows,[2.9,1.3,4.3,4.3,2.9],caption='表 10  相同 test 上的均值基线比较',left_cols=(0,))
p('基线恒定输出训练集各目标均值，测试 R² 约为零或为负。MAE 降低率按 1−模型 MAE/基线 MAE 计算。该比较确认模型学到了超出恒定均值的输入输出关系，尚不能替代与 MLP、树模型或只用压力的基准比较。')

# 12
page('11 真实值与预测值展示')
fig('03_test_parity','图 4  最终 test 的真实值与预测值  红色空心圆为 p_ocean=0 边界样本')
p('图中每个点对应一条完整工况样本，虚线表示理想预测。两组各目标的大多数点接近对角线，与表 9 的总体指标相符。红色空心点显示边界工况偏离较大，尤其是 ff。图中未删除误差较大记录，也未改变冻结分区。')
p('positive 测试 ff 最大绝对误差为 0.39339 N，发生于 (p_system,p_ocean)=(15,0) MPa，真实值 −0.721699 N、预测值 −0.328306 N。该记录属于 positive，不应与被排除的 negative (15,0) 混淆。negative 测试 ff 最大绝对误差为 0.27259 N，发生于 (60,0) MPa。')
p('f 的总体数值跨度大于 ff，高 R² 并不代表边界力误差足够小。positive 的最大 f 绝对误差为 8.95578 N，negative 为 13.58953 N。工程上是否可接受，仍需结合容许误差、工作安全裕度和量测精度单独判断。')
h('11.1 摩擦力输出的含义',2)
p('推理程序在逆标准化后计算 friction_prediction=ff_prediction+f_prediction，保证求和处于相同原单位。若需要评价总摩擦力，应逐样本将真实 ff+f 与预测 ff+f 比较；总摩擦力 MAE 或 RMSE 不能由两个目标的 MAE/RMSE 直接相加得到。本次冻结结果的正式指标维持 ff、f、q 三目标定义。')

# 13
page('12 误差分布与系统偏差')
fig('04_test_residuals','图 5  最终 test 残差随 p_ocean 的分布  残差定义为预测值减真实值')
rows=[]
for g in ('positive','negative'):
    for t in ('ff','f','q'):
        v=A['groups'][g]['test_error_distribution'][t]
        rows.append([g,t,num(v['bias']),num(v['median_abs']),num(v['p95_abs']),num(v['max_abs'])])
table(['分组','目标','平均残差','绝对误差中位数','绝对误差 P95','最大绝对误差'],rows,[2.7,1.1,2.8,3.6,3.3,3.2],
      caption='表 11  冻结测试预测的描述性误差统计',left_cols=(0,))
p('P95 是当前有限测试样本绝对误差的 95% 分位数，不是置信区间，也不是未来误差上界。negative f 的平均残差约 +1.04152 N，说明预测整体更偏向数值较大的一侧；对于负值 f，这对应多数预测较接近零。positive q 的平均残差约 +0.00416 mm³/s。')
p('局部偏差与极端误差需要与总体 MAE 同时解读。以上统计从已冻结的预测 CSV 计算，用于结果说明；没有再次执行模型、选择超参数或修改测试数据。')

# 14
page('13 边界工况的专项分析')
h('13.1 p_ocean 为零时的误差集中',2)
rows=[]
for g in ('positive','negative'):
    for label,display in [('p_ocean_zero','0'),('p_ocean_positive','>0')]:
        v=A['groups'][g]['subgroups']['test_'+label]
        rows.append([g,display,v['n'],num(v['MAE']['ff']),num(v['MAE']['f']),num(v['MAE']['q'])])
table(['分组','p_ocean','样本数','ff MAE  N','f MAE  N','q MAE  mm³/s'],rows,[2.7,2.0,1.8,3.1,3.1,4.0],
      caption='表 12  固定 test 中边界与其余工况的描述性对比',left_cols=(0,))
p('positive 边界测试仅 5 条，ff MAE=0.26960 N，明显高于其余 148 条的 0.01992 N；negative 边界测试仅 2 条，ff MAE=0.24458 N，高于其余 147 条的 0.01224 N。边界样本很少，不能用这些均值给出稳定的总体边界性能估计，但误差集中现象已经清晰。')
h('13.2 验证与测试难度存在差异',2)
p('negative 的两条边界 validation 为 (35,0)、(45,0) MPa，其 ff MAE=0.54110 N、f MAE=25.61614 N；test 边界为 (60,0)、(65,0) MPa，其 ff MAE=0.24458 N、f MAE=12.21784 N。验证难点与测试难点并不完全相同，因此不能用较好的总体 test 结果否定验证集中的边界问题。')
table(['工况  MPa','分区','ff 真实  N','ff 预测  N','绝对误差  N'],[
    ['positive 15 0','test','−0.721699','−0.328306','0.393393'],
    ['positive 65 0','test','−0.813880','−0.472661','0.341219'],
    ['negative 35 0','validation','1.871979','1.261168','0.610811'],
    ['negative 60 0','test','1.808812','1.536224','0.272588']
],[4.3,2.7,3.3,3.3,3.1],caption='表 13  具有代表性的已保存边界预测',left_cols=(0,))
h('13.3 结论与后续验证方向',2)
p('当前模型在多数已覆盖工况表现良好，边界跳变规律的学习仍不充分。边界现象可作为后续补充数据和独立压力留出实验的研究问题；本次不对 frozen test 做重加权、删点、边界补偿或继续调参。若方案改变，需要新的开发与独立评价协议。')

# 15
page('14 方案优劣与替代方法比较')
h('14.1 方法能力与代价',2)
table(['方案','主要优势','主要代价或限制','本项目证据'],[
    ['训练均值基线','无需学习输入关系  易复核','无法表达工况变化','已有同 test 实测'],
    ['当前 Transformer','共享节点映射  显式节点交互  融合压力','注意力成本随节点数平方增长  小样本易波动','已有冻结实测'],
    ['展平 MLP','实现直接  易构建回归基准','依赖固定列位置  缺少显式节点交互结构','未实施 方法分析'],
    ['节点汇总加树模型','可构建低维表格方案','汇总可能丢失分布  结果依赖特征工程','未实施 方法分析'],
    ['图神经网络','可利用已知邻接与几何关系','需要可靠节点对应与拓扑','未实施 方法分析']
],[2.8,4.6,5.5,3.8],caption='表 14  方法比较及证据等级',left_cols=(0,1,2,3))
p('Transformer 本身也是神经网络。与“神经网络”比较时，应明确基准为展平 MLP、节点 MLP 或其他具体结构。现有结果只支持当前 Transformer 优于均值基线，不能宣称其优于尚未在相同划分上实施的 MLP、树模型或图网络。')
h('14.2 当前方案的主要优点',2)
p('数据接口保留节点三元组关系，原值、清单、缩放和权重均可追溯；双模型避免未经确认的跨组混合。模型规模为每组 75843 个参数，流程能在现有 GPU 环境运行；标准化、验证早停、独立测试及无目标列推理形成了较完整的实验闭环。')
h('14.3 当前方案的主要局限',2)
p('每组样本量约千条，665 个输入和深层表示仍可能利用特定数据生成规律。无位置编码与均值池化不显式建模拓扑；压力后融合使注意力交互缺少直接压力条件化。所有节点平均汇聚也可能稀释少量关键区域的影响，这属于结构层面的可能原因，尚未通过消融验证。')
p('尚无只用压力、只用节点、去掉注意力或两组共享模型的对照，无法量化各分支的独立贡献。当前没有不确定性估计、物理约束或工况覆盖检查，也没有实测端到端部署加速。')

# 16
page('15 代码可靠性与可复现性评估')
h('15.1 已有验收与本次只读核验',2)
p('已有合成流程验收覆盖正负两组的形状检查、梯度、权重独立性、节点重排一致性、短程训练、checkpoint 重载、分目标指标、无目标列推理和泄漏拦截。真实数据审计记录包含原始哈希不变、行数守恒、节点抽查及关联组不跨分区。')
p('本次报告核对 v2 数据包所列文件哈希、run_004 两组权重哈希及元信息指纹；保存预测 CSV 的 sample_id 与固定 validation/test 集合相符，重新计算三目标指标与冻结 JSON 一致，缩放拟合 ID 仅来自 train。报告制作前后这些输入文件哈希保持不变，没有运行真实训练或重复 test 入口。')
h('15.2 代码中仍需维护的问题',2)
table(['发现','影响','建议'],[
    ['train_models.py 的 --data 默认仍为 v1','省略参数会与最终 v2 方案不一致','复现时显式 --data v2'],
    ['通用默认配置 lr=0.001','直接 both 训练不等于最终 positive 配置','两组分别指定最终 JSON'],
    ['根 README 含训练前历史说明','读者可能误判实验完成状态','以冻结元信息为准  后续同步文档'],
    ['未保存优化器状态与完整训练 RNG','当前版本不支持精确断点续训','需要时新增续训协议和状态存档'],
    ['未记录训练时源码快照哈希','现有权重指纹不能完整证明历史源码','后续保存 commit 或源文件清单']
],[6.0,5.4,5.3],caption='表 15  现有代码的可复现性与维护限制',left_cols=(0,1,2))
p('这些问题不改变已保存冻结成绩，但会影响新使用者准确复现最终流程。报告中的示例使用显式数据路径、配置和新输出目录；本次报告制作没有修改模型、训练入口或历史实验文件。')
h('15.3 冻结不等于完整实验隔离',2)
p('代码的重复评价拦截依赖已有结果文件及指纹，是工程保护而非不可绕过的统计隔离。后续开发应继续保持独立评价纪律；既有 frozen test 不应用于选择新损失、边界补偿、结构或异常阈值。若需要进一步研究，先约定新的开发和外部评价数据。')

# 17
page('16 局限 后续工作与结题结论')
h('16.1 物理与数据层面的未决事项',2)
p('x、y、cpress 的确切定义和单位、节点编号与物理点的对应、变形前后状态及新工况输入可获得性仍需来源方确认。源文件不能核验完整求解器残差与终止判据，潜在共同来源关联也不一定被完全记录。训练与高 R² 不足以确认这些物理事实。')
p('三条暂时排除记录不在最终适用数据域内；不能用现有成绩判断这三条工况的可预测性。已发现相同节点块的 11 条记录被分组隔离，但数值不完全相同的共同仿真来源仍可能有关联。当前压力分层随机留出无法证明范围外可靠性。')
h('16.2 后续工作优先级',2)
p('首先补齐变量字典、节点对应和源求解记录，明确工程允许误差；随后在预先约定的数据协议下补充只用压力、展平 MLP 以及节点汇总等基准，并进行节点分支、压力分支与注意力消融。所有比较应保持相同开发清单与评价单位，避免重抽测试集获得表面提升。')
p('对于未见压力等级，另设计完整 p_system 留出实验；对于更深海域或不同材料、几何、速度，需有对应范围的新独立数据。边界工况应获得足够样本并单列评价。多随机种子统计、不确定性估计和输入覆盖检测，可在已有基线与物理定义稳定后继续开展。')
h('16.3 结题结论',2)
p('本项目已交付从原始数据到独立冻结测试和新 CSV 推理的完整数值建模流程。最终 v2 实验使用 2090 条工况记录，positive、negative 分别训练同结构独立参数的两层 Transformer，输出 ff、f、q，原尺度测试 MAE 相对均值基线降低 91.09% 至 97.30%。')
p('positive 测试三目标 R² 为 0.9754、0.9960、0.9910，negative 为 0.9930、0.9990、0.9969，证明了现有数据条件下较强的回归能力。结果同时揭示 p_ocean=0 边界样本的显著误差，方案尚不能作为外推保证或实际工程验收依据。结题成果是有清晰来源、固定评价和明确适用边界的实验方案，为后续物理验证及方法比较提供可复用基础。')

# 18
page('附录 A 最终方案运行与交付物')
h('A1 复现最终配置的命令',2)
p('以下命令用于今后在项目根目录按最终配置创建新的训练实验。已有 run_004 及测试文件应保留。命令显式使用 pytorch2、v2 数据与分组配置，不依赖旧默认值。')
code("python scripts/train_models.py `\n  --device cuda --data outputs/training_data/v2 --check-data\n\npython scripts/train_models.py `\n  --device cuda --data outputs/training_data/v2 `\n  --config scripts/training_positive.json --group positive `\n  --output outputs/training/run_report_reproduction\n\npython scripts/train_models.py `\n  --device cuda --data outputs/training_data/v2 `\n  --config scripts/training_negative.json --group negative `\n  --output outputs/training/run_report_reproduction")
p('新实验完成后先使用 validation 评估。run_004 的 test 已冻结，不再执行其评价命令；既有 test_metrics.json 与 test_predictions.csv 直接用于结果展示。重训可能存在 GPU 数值差异，且不应借重新训练继续针对旧 test 选模型。')
h('A2 无目标列推理',2)
code("python scripts/predict.py `\n  --device cuda --run outputs/training/run_004/positive `\n  --input new_positive.csv --output outputs/predictions/positive_new.csv")
p('new_positive.csv 应包含两个压力及全部 663 个节点输入，ff/f/q 可以不存在；若包含 group，必须与模型分组相符。推理输出 ff_prediction、f_prediction、q_prediction 及 friction_prediction，输入和已有输出均不覆盖。')
h('A3 核心交付物',2)
table(['交付物','项目相对位置'],[
    ['最终模型与预测','outputs/training/run_004/positive 及 negative'],
    ['最终实验数据和清单','outputs/training_data/v2'],
    ['数据清洗证据','outputs/data_preparation/evidence_review'],
    ['Word 报告及 Markdown','outputs/final_report'],
    ['报告图表与核验统计','outputs/final_report/figures 及 analysis.json'],
    ['报告可重建脚本','scripts/report_analysis.py 及 build_final_report.py']
],[5.6,11.1],caption='表 A1  交付物索引',left_cols=(0,1))

# 19
page('附录 B 证据索引与参考文献')
h('B1 项目证据索引',2)
p('正文 [R1] 至 [R3] 为方法文献，项目实验数字由下列本地证据支持。路径均相对于项目根目录。图表来自已保存数据与预测文件；没有生成新的模型测试成绩。')
table(['正文内容','主要证据'],[
    ['输入与模型结构','models/data.py  models/seal_transformer.py'],
    ['异常和关联处置','scripts/data_exclusions.json  scripts/finalize_cleaning_review.py'],
    ['清洗与来源核对','outputs/data_preparation/evidence_review/cleaning_verification.json'],
    ['实际数据及状态','outputs/training_data/v2/dataset_metadata.json  split_manifest.csv'],
    ['训练策略与配置','scripts/train_models.py  training_positive.json  training_negative.json'],
    ['开发对比','outputs/training/run_001 至 run_003 各组 metadata.json'],
    ['最终训练与验证','outputs/training/run_004 各组 metadata.json  history.csv'],
    ['冻结测试及绘图','outputs/training/run_004 各组 test_metrics.json  test_predictions.csv'],
    ['推理与历史验收','scripts/predict.py  outputs/model_checks/pipeline_checks.json'],
    ['本次报告核验','outputs/final_report/analysis.json']
],[4.2,12.5],caption='表 B1  报告内容与代码及产物对照',left_cols=(0,1))
p('positive best.pt 的 SHA256 为\n'+A['groups']['positive']['checkpoint_sha256'])
p('negative best.pt 的 SHA256 为\n'+A['groups']['negative']['checkpoint_sha256'])
p('完整的本次读取文件 SHA256、指标一致性检查和分组统计保存在 analysis.json。当前源码的哈希记录用于本次报告复核，不等同于训练时源码快照。')
h('B2 方法参考文献',2)
p('[R1] Vaswani A, et al. Attention Is All You Need. 2017. https://arxiv.org/abs/1706.03762')
p('[R2] Zaheer M, et al. Deep Sets. 2017. https://arxiv.org/abs/1703.06114')
p('[R3] Loshchilov I, Hutter F. Decoupled Weight Decay Regularization. ICLR 2019. https://arxiv.org/abs/1711.05101')
p('文献说明注意力、集合不变性与 AdamW 的方法依据；本项目没有采用预训练权重，也不将上述文献中的其他任务成绩当作本项目性能证据。')

D.save(DOCX)
MD.write_text('# 深海柔性密封性能预测项目结题报告\n\n'+'\n'.join(markdown),encoding='utf-8')
print(DOCX)
print(MD)
print('Paragraphs',len(D.paragraphs),'Tables',len(D.tables),'Figures',len(D.inline_shapes))
